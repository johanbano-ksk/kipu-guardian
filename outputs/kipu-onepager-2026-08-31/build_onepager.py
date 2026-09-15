from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_DIR = Path(r"C:\Users\johan.bano_kushkipag\Documents\Dev\kipu-alert-reviewer\outputs\kipu-onepager-2026-08-31")
OUT_PATH = OUT_DIR / "KIPU_Guardian_OnePager_Roadmap_2026-08-31.docx"

# standard_business_brief + named brand overrides for Kushki and one-page density.
NAVY = "023365"
NAVY_DEEP = "012746"
INK = "082F55"
MINT = "00E6B2"
MINT_SOFT = "DFFBF4"
BLUE = "1E65AE"
SKY_SOFT = "EEF6FC"
LINE = "D5E4F2"
MUTED = "677784"
WHITE = "FFFFFF"
LIGHT = "F6F9FC"
GREEN = "16835B"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, attrs in edges.items():
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key, value in attrs.items():
            element.set(qn(f"w:{key}"), str(value))


def set_table_geometry(table, widths_in, indent_dxa=120):
    widths_dxa = [int(round(width * 1440)) for width in widths_in]
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths_in[idx])
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_run(run, *, size=10.2, bold=False, color=INK, italic=False, font="Calibri"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)


def format_paragraph(p, *, before=0, after=4, line=1.05, keep_next=False):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.keep_with_next = keep_next


def add_text(p, text, **kwargs):
    run = p.add_run(text)
    set_run(run, **kwargs)
    return run


def set_paragraph_bottom_border(p, color=MINT, size=18, space=5):
    p_pr = p._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.05

    for style_name, size, before, after in (
        ("Heading 1", 15.5, 9, 4),
        ("Heading 2", 12.5, 7, 3),
        ("Heading 3", 11.3, 5, 2),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(NAVY)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    bullet = styles["List Bullet"]
    bullet.font.name = "Calibri"
    bullet.font.size = Pt(9.8)
    bullet.font.color.rgb = rgb(INK)
    bullet.paragraph_format.left_indent = Inches(0.42)
    bullet.paragraph_format.first_line_indent = Inches(-0.20)
    bullet.paragraph_format.space_after = Pt(2.5)
    bullet.paragraph_format.line_spacing = 1.03


def add_hyperlink(paragraph, text, url, color=BLUE):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    size_el = OxmlElement("w:sz")
    size_el.set(qn("w:val"), "18")
    r_pr.extend([r_fonts, color_el, underline, size_el])
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.extend([r_pr, text_el])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_section_heading(doc, text):
    p = doc.add_paragraph(style="Heading 2")
    add_text(p, text, size=12.5, bold=True, color=NAVY)
    return p


doc = Document()
configure_styles(doc)
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
# Named override: compact one-page geometry while preserving a balanced letter format.
section.top_margin = Inches(0.56)
section.bottom_margin = Inches(0.52)
section.left_margin = Inches(0.65)
section.right_margin = Inches(0.65)
section.header_distance = Inches(0.28)
section.footer_distance = Inches(0.28)
content_width = 7.2

# Running header/footer.
header_p = section.header.paragraphs[0]
header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
format_paragraph(header_p, after=0, line=1.0)
add_text(header_p, "KIPU GUARDIAN", size=8.5, bold=True, color=NAVY)
add_text(header_p, "  |  PAYMENTS INTELLIGENCE", size=8.5, color=MUTED)

footer_p = section.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
format_paragraph(footer_p, after=0, line=1.0)
add_text(footer_p, "Uso interno  |  Estado al 31 ago 2026", size=8, color=MUTED)

# Brand band.
band = doc.add_table(rows=1, cols=2)
set_table_geometry(band, [4.4, 2.8], indent_dxa=0)
for cell in band.rows[0].cells:
    set_cell_shading(cell, NAVY)
    set_cell_border(
        cell,
        top={"val": "nil"}, bottom={"val": "nil"},
        start={"val": "nil"}, end={"val": "nil"},
    )
    set_cell_margins(cell, top=105, bottom=105, start=160, end=160)
p = band.cell(0, 0).paragraphs[0]
format_paragraph(p, after=0, line=1.0)
add_text(p, "KUSHKI", size=15, bold=True, color=WHITE)
add_text(p, "  +", size=15, bold=True, color=MINT)
p = band.cell(0, 1).paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
format_paragraph(p, after=0, line=1.0)
add_text(p, "ONE-PAGER DE ROADMAP", size=9, bold=True, color=MINT)

# Memo masthead.
title = doc.add_paragraph()
format_paragraph(title, before=7, after=1, line=1.0, keep_next=True)
add_text(title, "KIPU Guardian", size=24, bold=True, color=NAVY_DEEP)
subtitle = doc.add_paragraph()
format_paragraph(subtitle, after=6, line=1.0, keep_next=True)
add_text(
    subtitle,
    "Revisión automatizada y trazable de alertas críticas KIPU 1.0",
    size=11.5,
    color=MUTED,
)
set_paragraph_bottom_border(subtitle)

# Metric strip: compact factual status summary.
metrics = doc.add_table(rows=2, cols=4)
set_table_geometry(metrics, [1.55, 1.85, 2.10, 1.70], indent_dxa=0)
metric_data = [
    ("ESTADO", "LIVE EN DEV", GREEN),
    ("HORIZONTE", "Q3 - Q4 2026", NAVY),
    ("POLÍTICA", "CRÍTICAS · DETERMINÍSTICA", NAVY),
    ("IMPACTO META", "15 H / SEMANA", NAVY),
]
for idx, (label, value, value_color) in enumerate(metric_data):
    top = metrics.cell(0, idx)
    bottom = metrics.cell(1, idx)
    set_cell_shading(top, SKY_SOFT)
    set_cell_shading(bottom, LIGHT)
    for cell in (top, bottom):
        set_cell_border(
            cell,
            top={"val": "single", "sz": 5, "color": LINE},
            bottom={"val": "single", "sz": 5, "color": LINE},
            start={"val": "single", "sz": 5, "color": LINE},
            end={"val": "single", "sz": 5, "color": LINE},
        )
        set_cell_margins(cell, top=70, bottom=70, start=95, end=95)
    p = top.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p, after=0, line=1.0)
    add_text(p, label, size=7.7, bold=True, color=MUTED)
    p = bottom.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p, after=0, line=1.0)
    add_text(p, value, size=9.3, bold=True, color=value_color)

# Lead callout.
callout = doc.add_table(rows=1, cols=1)
set_table_geometry(callout, [content_width], indent_dxa=0)
cell = callout.cell(0, 0)
set_cell_shading(cell, MINT_SOFT)
set_cell_border(
    cell,
    start={"val": "single", "sz": 24, "color": MINT},
    top={"val": "nil"}, bottom={"val": "nil"}, end={"val": "nil"},
)
set_cell_margins(cell, top=105, bottom=105, start=150, end=140)
p = cell.paragraphs[0]
format_paragraph(p, after=0, line=1.08)
add_text(p, "PROPÓSITO  ", size=8.2, bold=True, color=NAVY)
add_text(
    p,
    "Automatizar la revisión de alertas, mostrar sólo eventos críticos con evidencia estructurada y reducir el tiempo operativo sin ocultar falsos positivos ni perder trazabilidad.",
    size=10.2,
    bold=True,
    color=INK,
)

add_section_heading(doc, "Qué está operativo hoy")
capabilities = doc.add_table(rows=2, cols=2)
set_table_geometry(capabilities, [3.6, 3.6], indent_dxa=0)
cap_data = [
    ("01  Ingesta y resiliencia", "EventBridge > SQS > Lambda. Reintentos y desacoplamiento del procesamiento."),
    ("02  Decisión consistente", "Filtro crítico KIPU 1.0 basado en métricas estructuradas; sin dependencia de búsquedas externas."),
    ("03  Trazabilidad", "Archivo de ocurrencias e idempotencia en DynamoDB para evitar duplicados y facilitar auditoría."),
    ("04  Operación visible", "Dashboard con MID, país, fecha/hora generada en KIPU y extracción manual bajo demanda."),
]
for idx, (heading, body) in enumerate(cap_data):
    row, col = divmod(idx, 2)
    cell = capabilities.cell(row, col)
    set_cell_shading(cell, WHITE if row == 0 else LIGHT)
    set_cell_border(
        cell,
        top={"val": "single", "sz": 5, "color": LINE},
        bottom={"val": "single", "sz": 5, "color": LINE},
        start={"val": "single", "sz": 5, "color": LINE},
        end={"val": "single", "sz": 5, "color": LINE},
    )
    set_cell_margins(cell, top=95, bottom=95, start=125, end=125)
    p = cell.paragraphs[0]
    format_paragraph(p, after=2, line=1.0, keep_next=True)
    add_text(p, heading, size=9.5, bold=True, color=NAVY)
    p = cell.add_paragraph()
    format_paragraph(p, after=0, line=1.05)
    add_text(p, body, size=9.2, color=INK)

add_section_heading(doc, "Roadmap exclusivo del proyecto")
roadmap = doc.add_table(rows=1, cols=3)
set_table_geometry(roadmap, [2.35, 2.42, 2.43], indent_dxa=0)
roadmap_items = [
    (
        "AHORA · MVP",
        "Completado",
        "Operación en ambiente dev, repositorio privado, política determinística, dashboard y ejecución manual.",
        MINT_SOFT,
    ),
    (
        "CIERRE Q3 · ESTABILIZAR",
        "Siguiente hito",
        "Conciliar KIPU/Slack/dashboard por país, medir precisión y falsos positivos, y endurecer observabilidad y accesos.",
        SKY_SOFT,
    ),
    (
        "Q4 · ESCALAR",
        "Evolución controlada",
        "Integrar el contrato común de agentes, validar producción y evaluar IA sólo si mejora el benchmark con control humano.",
        LIGHT,
    ),
]
for idx, (phase, status, body, fill) in enumerate(roadmap_items):
    cell = roadmap.cell(0, idx)
    set_cell_shading(cell, fill)
    set_cell_border(
        cell,
        top={"val": "single", "sz": 8, "color": MINT if idx == 0 else LINE},
        bottom={"val": "single", "sz": 5, "color": LINE},
        start={"val": "single", "sz": 5, "color": LINE},
        end={"val": "single", "sz": 5, "color": LINE},
    )
    set_cell_margins(cell, top=105, bottom=105, start=120, end=120)
    p = cell.paragraphs[0]
    format_paragraph(p, after=3, line=1.0, keep_next=True)
    add_text(p, phase, size=8.1, bold=True, color=NAVY)
    p = cell.add_paragraph()
    format_paragraph(p, after=3, line=1.0, keep_next=True)
    add_text(p, status, size=10.2, bold=True, color=GREEN if idx == 0 else BLUE)
    p = cell.add_paragraph()
    format_paragraph(p, after=0, line=1.05)
    add_text(p, body, size=8.9, color=INK)

add_section_heading(doc, "Cómo sabremos que funciona")
kpi_table = doc.add_table(rows=2, cols=3)
set_table_geometry(kpi_table, [2.4, 2.4, 2.4], indent_dxa=0)
kpis = [
    ("COBERTURA", "% alertas KIPU procesadas"),
    ("CALIDAD", "Precisión y tasa de falsos positivos"),
    ("VELOCIDAD", "Latencia alerta > dashboard"),
    ("CONFIABILIDAD", "Duplicados, errores y disponibilidad"),
    ("IMPACTO", "Horas liberadas verificadas"),
    ("ADOPCIÓN", "Revisiones y ejecuciones manuales"),
]
for idx, (label, value) in enumerate(kpis):
    row, col = divmod(idx, 3)
    cell = kpi_table.cell(row, col)
    set_cell_shading(cell, WHITE)
    set_cell_border(
        cell,
        bottom={"val": "single", "sz": 5, "color": LINE},
        top={"val": "nil"}, start={"val": "nil"}, end={"val": "nil"},
    )
    set_cell_margins(cell, top=55, bottom=55, start=80, end=80)
    p = cell.paragraphs[0]
    format_paragraph(p, after=1, line=1.0)
    add_text(p, label, size=7.5, bold=True, color=MUTED)
    p = cell.add_paragraph()
    format_paragraph(p, after=0, line=1.0)
    add_text(p, value, size=8.8, bold=True, color=NAVY)

# Decision and evidence note.
decision = doc.add_paragraph()
format_paragraph(decision, before=5, after=2, line=1.05)
add_text(decision, "Gate de salida a producción: ", size=9.2, bold=True, color=NAVY)
add_text(
    decision,
    "no activar IA ni declarar el ahorro de 15 h/semana como resultado hasta contar con datos reconciliados, benchmark de precisión y controles operativos acordados.",
    size=9.2,
    color=INK,
)

links = doc.add_paragraph()
format_paragraph(links, after=0, line=1.0)
add_text(links, "Activos: ", size=8.8, bold=True, color=MUTED)
add_hyperlink(links, "Dashboard", "https://kipu-alert-reviewer-dashboard.kushkipagos-9534.chatgpt.site/")
add_text(links, "  ·  ", size=8.8, color=MUTED)
add_hyperlink(links, "Repositorio privado", "https://github.com/johanbano-ksk/kipu-alert-reviewer")

doc.core_properties.title = "KIPU Guardian - One-pager de roadmap"
doc.core_properties.subject = "Roadmap exclusivo del proyecto KIPU Guardian"
doc.core_properties.keywords = "KIPU, alertas, roadmap, AWS, dashboard, Kushki"
doc.core_properties.author = "Payments Intelligence"
doc.core_properties.comments = "Estado del proyecto al 31 de agosto de 2026."

OUT_DIR.mkdir(parents=True, exist_ok=True)
doc.save(OUT_PATH)
print(OUT_PATH)
